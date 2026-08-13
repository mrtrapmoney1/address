#!/usr/bin/env python3
"""
Check-Workpaper.py  -  scan a finished workpaper for the "scramble" defect.

The old script sometimes jammed the code + row + flag into a single cell
("236652 0", "OK 0", "FALLBACK - row 97027 (unvalidated) 0"). This scans the
result block and flags any row where that happened, so you can catch a bad run
before you trust it.

Usage:
    python3 tests/Check-Workpaper.py <workbook.xlsx> [--sheet Sheet1]
                                     [--header-row 4] [--code-col "NE City Code"]

It finds the result block by header name, then checks every data row:
  * a flag cell must not hold a trailing jammed number ("<text> <number>")
  * the code cell must be blank or an integer 0..999
Exit code is 0 when clean, 1 when any scrambled/invalid row is found.

Requires openpyxl (pip install openpyxl).
"""
import argparse, re, sys

def col_index(ws, header_row, name):
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        if v is not None and str(v).strip() == name:
            return c
    return None

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--sheet", default=None)
    ap.add_argument("--header-row", type=int, default=4)
    ap.add_argument("--code-col", default="NE City Code")
    ap.add_argument("--flag-col", default="Match Flag")
    args = ap.parse_args()

    try:
        import openpyxl
    except ImportError:
        print("ERROR: openpyxl is required (pip install openpyxl)"); return 2

    wb = openpyxl.load_workbook(args.workbook, data_only=True)
    ws = wb[args.sheet] if args.sheet else wb[wb.sheetnames[0]]
    hr = args.header_row

    code_c = col_index(ws, hr, args.code_col)
    flag_c = col_index(ws, hr, args.flag_col)   # may be absent (code-only output)
    if code_c is None:
        print(f"ERROR: could not find a '{args.code_col}' header on row {hr} of '{ws.title}'.")
        return 2

    jam = re.compile(r".+\s-?\d+$")
    scrambled, bad_code = [], []
    for r in range(hr + 1, ws.max_row + 1):
        code = ws.cell(r, code_c).value
        flag = ws.cell(r, flag_c).value if flag_c else None
        # scramble: a flag cell that ends in a jammed number
        if isinstance(flag, str) and jam.match(flag.strip()) and "REPAIRED" not in flag:
            scrambled.append((r, flag))
        # invalid code: present but not an integer 0..999
        if code is not None and str(code).strip() != "":
            try:
                n = float(code)
                if n < 0 or n > 999 or n != int(n):
                    bad_code.append((r, code))
            except (TypeError, ValueError):
                bad_code.append((r, code))

    print(f"Workbook : {args.workbook}")
    print(f"Sheet    : {ws.title}   header row {hr}")
    print(f"Code col : {args.code_col} (column {code_c})"
          + (f"   Flag col: {args.flag_col} (column {flag_c})" if flag_c else "   (no flag column - code only)"))
    print(f"Data rows: {ws.max_row - hr}")
    print("-" * 60)
    if not scrambled and not bad_code:
        print("CLEAN - no scrambled cells, every code is a valid 0..999.")
        return 0
    if scrambled:
        print(f"SCRAMBLED cells (code/row jammed into the flag): {len(scrambled)}")
        for r, f in scrambled[:15]:
            print(f"   row {r}: flag = {f!r}")
        if len(scrambled) > 15:
            print(f"   ... and {len(scrambled) - 15} more")
    if bad_code:
        print(f"INVALID codes (not an integer 0..999): {len(bad_code)}")
        for r, c in bad_code[:15]:
            print(f"   row {r}: code = {c!r}")
    return 1

if __name__ == "__main__":
    sys.exit(main())
